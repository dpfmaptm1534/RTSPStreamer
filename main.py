import csv
import json
import os
import queue
import re
import socket
import subprocess
import sys
import threading
import time
import tkinter as tk
from datetime import datetime
from tkinter import filedialog, messagebox, ttk
from urllib.parse import urlparse

import openpyxl
import psycopg2
import xlrd
from psycopg2 import sql

try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
except ImportError:
    DND_FILES = None
    TkinterDnD = None


if getattr(sys, "frozen", False):
    CURRENT_DIR = os.path.dirname(os.path.abspath(sys.executable))
else:
    CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))

FFMPEG_PATH = os.path.join(CURRENT_DIR, "ffmpeg.exe")
MEDIAMTX_PATH = os.path.join(CURRENT_DIR, "mediamtx.exe")
MEDIAMTX_CONFIG_PATH = os.path.join(CURRENT_DIR, "mediamtx.yml")
MEDIAMTX_LOG_PATH = os.path.join(CURRENT_DIR, "mediamtx.log")
FFMPEG_LOG_PATH = os.path.join(CURRENT_DIR, "ffmpeg.log")
DEFAULT_RTSP_URL = "rtsp://127.0.0.1:8554/live/stream"
RECENT_SETTINGS_PATH = os.path.join(CURRENT_DIR, "recent_settings.json")
CREATE_NO_WINDOW = 0x08000000 if os.name == "nt" else 0


class RTSPStreamerGUI:
    def __init__(self, root):
        self.root = root
        self.root.title("영상/웹캠 RTSP 송출기 + DB 입력")
        self.root.geometry("1050x760")
        self.root.minsize(900, 520)

        self.video_path = ""
        self.webcam_devices = []
        self.video_duration = 0.0
        self.schedule_rows = []
        self.ffmpeg_process = None
        self.preview_process = None
        self.preview_thread = None
        self.preview_image = None
        self.preview_queue = queue.Queue()
        self.preview_polling = False
        self.server_process = None
        self.server_log_file = None
        self.db_thread = None
        self.recent_settings = []
        self.is_streaming = False
        self.closing = False
        self.edit_entry = None
        self.active_cell_item = ""
        self.active_cell_column = "value"
        self.cell_border_parts = []

        self.create_widgets()
        self.root.protocol("WM_DELETE_WINDOW", self.close)
        self.start_internal_server()

    def create_widgets(self):
        action_frame = tk.Frame(self.root, bd=1, relief="raised", bg="#f3f4f6")
        action_frame.pack(side="bottom", fill="x", padx=0, pady=0)

        self.lbl_status = tk.Label(
            action_frame,
            text="상태: 대기 중",
            fg="#2563eb",
            bg="#f3f4f6",
            font=("맑은 고딕", 10, "bold"),
        )
        self.lbl_status.pack(side="left", padx=18, pady=10)

        self.btn_action = tk.Button(
            action_frame,
            text="송출 시작",
            bg="#16a34a",
            fg="white",
            font=("맑은 고딕", 11, "bold"),
            width=16,
            height=2,
            command=self.toggle_streaming,
            state="disabled",
        )
        self.btn_action.pack(side="right", padx=18, pady=8)

        scroll_container = tk.Frame(self.root)
        scroll_container.pack(side="top", fill="both", expand=True)

        self.main_canvas = tk.Canvas(scroll_container, highlightthickness=0)
        self.main_scrollbar = ttk.Scrollbar(
            scroll_container, orient="vertical", command=self.main_canvas.yview
        )
        self.main_canvas.configure(yscrollcommand=self.main_scrollbar.set)
        self.main_scrollbar.pack(side="right", fill="y")
        self.main_canvas.pack(side="left", fill="both", expand=True)

        self.main_content = tk.Frame(self.main_canvas)
        self.main_canvas_window = self.main_canvas.create_window(
            (0, 0), window=self.main_content, anchor="nw"
        )
        self.main_content.bind("<Configure>", self.update_main_scroll_region)
        self.main_canvas.bind("<Configure>", self.resize_main_canvas_window)
        self.main_canvas.bind("<Enter>", self.enable_main_mousewheel)
        self.main_canvas.bind("<Leave>", self.disable_main_mousewheel)

        top_frame = tk.Frame(self.main_content)
        top_frame.pack(fill="x", padx=14, pady=(12, 6))

        settings_frame = tk.Frame(top_frame)
        settings_frame.pack(fill="x", pady=(0, 8))

        tk.Button(settings_frame, text="설정 저장", width=10, command=self.save_settings_as).pack(
            side="left"
        )
        tk.Button(settings_frame, text="설정 불러오기", width=12, command=self.load_settings_from_file).pack(
            side="left", padx=(6, 0)
        )

        tk.Label(settings_frame, text="최근 설정").pack(side="right", padx=(8, 0))
        self.recent_settings_var = tk.StringVar(value="")
        self.combo_recent_settings = ttk.Combobox(
            settings_frame,
            width=42,
            textvariable=self.recent_settings_var,
            state="readonly",
            values=[],
        )
        self.combo_recent_settings.pack(side="right")
        self.combo_recent_settings.bind("<<ComboboxSelected>>", self.load_recent_setting)
        self.load_recent_settings()

        file_frame = tk.LabelFrame(top_frame, text=" 1. 입력 소스 선택 ", padx=10, pady=10)
        file_frame.pack(fill="x")
        self.source_frame = file_frame

        self.source_mode_var = tk.StringVar(value="file")
        tk.Radiobutton(
            file_frame,
            text="영상 파일",
            variable=self.source_mode_var,
            value="file",
            command=self.update_source_mode_ui,
        ).pack(side="left", padx=(0, 8))
        tk.Radiobutton(
            file_frame,
            text="웹캠",
            variable=self.source_mode_var,
            value="webcam",
            command=self.update_source_mode_ui,
        ).pack(side="left", padx=(0, 8))

        self.lbl_file = tk.Label(file_frame, text="선택된 파일이 없습니다.", fg="gray", anchor="w")
        self.lbl_file.pack(side="left", fill="x", expand=True, padx=5)

        self.lbl_duration = tk.Label(file_frame, text="", fg="#555555")
        self.lbl_duration.pack(side="left", padx=8)

        self.webcam_var = tk.StringVar(value="")
        self.combo_webcam = ttk.Combobox(
            file_frame,
            width=28,
            textvariable=self.webcam_var,
            state="readonly",
            values=[],
        )
        self.combo_webcam.bind("<<ComboboxSelected>>", self.on_webcam_selected)
        self.btn_refresh_webcam = tk.Button(
            file_frame, text="웹캠 새로고침", width=12, command=self.refresh_webcam_devices
        )

        self.btn_browse = tk.Button(file_frame, text="파일 선택", width=11, command=self.browse_file)
        self.btn_browse.pack(side="right", padx=5)

        self.preview_frame = tk.LabelFrame(top_frame, text=" 웹캠 미리보기 ", padx=10, pady=8)
        self.lbl_preview = tk.Label(
            self.preview_frame,
            text="웹캠을 선택하면 여기에 미리보기가 표시됩니다.",
            bg="black",
            fg="white",
            width=44,
            height=10,
            anchor="center",
        )
        self.lbl_preview.pack(side="left")
        self.lbl_preview_status = tk.Label(
            self.preview_frame,
            text="",
            fg="#555555",
            anchor="w",
            justify="left",
        )
        self.lbl_preview_status.pack(side="left", fill="x", expand=True, padx=12)
        self.update_source_mode_ui()

        url_frame = tk.LabelFrame(top_frame, text=" 2. RTSP 송출 주소 ", padx=10, pady=10)
        url_frame.pack(fill="x", pady=(8, 0))

        self.entry_url = tk.Entry(url_frame)
        self.entry_url.insert(0, DEFAULT_RTSP_URL)
        self.entry_url.pack(fill="x", padx=5)

        self.lbl_server = tk.Label(
            top_frame, text="RTSP 서버: 시작 중...", fg="#d97706", anchor="w"
        )
        self.lbl_server.pack(fill="x", padx=8, pady=(8, 0))

        mid_frame = tk.Frame(self.main_content)
        mid_frame.pack(fill="both", expand=True, padx=14, pady=6)

        db_frame = tk.LabelFrame(mid_frame, text=" 3. DB INSERT 설정 ", padx=10, pady=8)
        db_frame.pack(fill="x")

        self.db_enabled_var = tk.BooleanVar(value=True)
        tk.Checkbutton(
            db_frame,
            text="송출 중 시간표에 맞춰 DB INSERT 사용",
            variable=self.db_enabled_var,
        ).grid(row=0, column=0, columnspan=4, sticky="w", pady=(0, 6))

        tk.Label(db_frame, text="DB 명 (필수)").grid(row=1, column=0, sticky="w")
        self.entry_db_name = tk.Entry(db_frame)
        self.entry_db_name.insert(0, "MFLOW")
        self.entry_db_name.grid(row=1, column=1, sticky="ew", padx=5)

        tk.Label(db_frame, text="DB 사용자 명 (필수)").grid(row=1, column=2, sticky="e")
        self.entry_db_user = tk.Entry(db_frame)
        self.entry_db_user.insert(0, "postgres")
        self.entry_db_user.grid(row=1, column=3, sticky="ew", padx=(5, 0))

        tk.Label(db_frame, text="패스워드(암호) (필수)").grid(row=2, column=0, sticky="w", pady=(6, 0))
        self.entry_db_password = tk.Entry(db_frame, show="*")
        self.entry_db_password.grid(row=2, column=1, sticky="ew", padx=5, pady=(6, 0))

        tk.Label(db_frame, text="접속주소(host) (필수)").grid(row=2, column=2, sticky="e", pady=(6, 0))
        self.entry_db_host = tk.Entry(db_frame)
        self.entry_db_host.insert(0, "localhost")
        self.entry_db_host.grid(row=2, column=3, sticky="ew", padx=(5, 0), pady=(6, 0))

        tk.Label(db_frame, text="접속포트(port) (필수)").grid(row=3, column=0, sticky="w", pady=(6, 0))
        self.entry_db_port = tk.Entry(db_frame)
        self.entry_db_port.insert(0, "5432")
        self.entry_db_port.grid(row=3, column=1, sticky="ew", padx=5, pady=(6, 0))

        tk.Button(db_frame, text="DB 연결 테스트", command=self.test_db_connection).grid(
            row=3, column=3, sticky="e", pady=(6, 0)
        )

        tk.Label(db_frame, text="스키마").grid(row=4, column=0, sticky="w", pady=(6, 0))
        self.entry_schema = tk.Entry(db_frame)
        self.entry_schema.insert(0, "MEASURE")
        self.entry_schema.grid(row=4, column=1, sticky="ew", padx=5, pady=(6, 0))

        tk.Label(db_frame, text="테이블").grid(row=4, column=2, sticky="e", pady=(6, 0))
        self.entry_table = tk.Entry(db_frame)
        self.entry_table.insert(0, "water_level_log")
        self.entry_table.grid(row=4, column=3, sticky="ew", padx=(5, 0), pady=(6, 0))

        tk.Label(db_frame, text="현재시간 컬럼").grid(row=5, column=0, sticky="w", pady=(6, 0))
        self.entry_time_col = tk.Entry(db_frame, width=18)
        self.entry_time_col.insert(0, "created_at")
        self.entry_time_col.grid(row=5, column=1, sticky="ew", padx=5, pady=(6, 0))

        tk.Label(db_frame, text="수위값 컬럼").grid(row=5, column=2, sticky="e", pady=(6, 0))
        self.entry_value_col = tk.Entry(db_frame, width=18)
        self.entry_value_col.insert(0, "level_value")
        self.entry_value_col.grid(row=5, column=3, sticky="ew", padx=(5, 0), pady=(6, 0))

        db_frame.columnconfigure(1, weight=1)
        db_frame.columnconfigure(3, weight=1)

        table_frame = tk.LabelFrame(mid_frame, text=" 4. 영상 시간별 수위값 표 ", padx=10, pady=8)
        self.table_frame = table_frame
        table_frame.pack(fill="both", expand=True, pady=(8, 0))

        toolbar = tk.Frame(table_frame)
        toolbar.pack(fill="x", pady=(0, 6))

        tk.Label(toolbar, text="간격").pack(side="left")
        self.interval_value_var = tk.StringVar(value="5")
        self.spin_interval = tk.Spinbox(
            toolbar, from_=1, to=3600, width=6, textvariable=self.interval_value_var
        )
        self.spin_interval.pack(side="left", padx=(5, 4))

        self.interval_unit_var = tk.StringVar(value="초")
        ttk.Combobox(
            toolbar,
            width=6,
            textvariable=self.interval_unit_var,
            values=["초", "분"],
            state="readonly",
        ).pack(side="left")

        tk.Button(toolbar, text="표 다시 만들기", command=self.rebuild_schedule_from_ui).pack(
            side="left", padx=8
        )
        tk.Button(toolbar, text="행 추가", command=self.add_schedule_row).pack(side="left")
        tk.Button(toolbar, text="선택 행 삭제", command=self.delete_selected_rows).pack(
            side="left", padx=(6, 0)
        )
        tk.Button(toolbar, text="수위값 비우기", command=self.clear_values).pack(side="left")
        tk.Button(toolbar, text="수위파일 불러오기", command=self.import_level_file).pack(
            side="left", padx=(8, 0)
        )

        self.skip_file_first_row_var = tk.BooleanVar(value=True)
        tk.Checkbutton(
            toolbar,
            text="파일 첫 행 건너뛰기",
            variable=self.skip_file_first_row_var,
        ).pack(side="left", padx=(8, 0))

        self.skip_table_first_row_var = tk.BooleanVar(value=False)
        tk.Checkbutton(
            toolbar,
            text="표 첫 행 건너뛰기",
            variable=self.skip_table_first_row_var,
        ).pack(side="left", padx=(8, 0))

        self.lbl_paste_help = tk.Label(
            toolbar,
            text="엑셀/CSV 드래그 또는 불러오기: 선택 행부터 수위값 채움",
            fg="#555555",
        )
        self.lbl_paste_help.pack(side="left", padx=12)

        self.lbl_db_status = tk.Label(toolbar, text="DB INSERT: 대기", fg="#555555")
        self.lbl_db_status.pack(side="right")

        tree_box = tk.Frame(table_frame)
        tree_box.pack(fill="both", expand=True)

        self.tree = ttk.Treeview(
            tree_box,
            columns=("time", "value"),
            show="headings",
            selectmode="extended",
            height=14,
        )
        self.tree.heading("time", text="영상 시간")
        self.tree.heading("value", text="수위값")
        self.tree.column("time", width=140, anchor="center", stretch=False)
        self.tree.column("value", width=780, anchor="w", stretch=True)

        style = ttk.Style()
        style.map(
            "Treeview",
            background=[("selected", "white")],
            foreground=[("selected", "black")],
        )

        self.cell_border_parts = [tk.Frame(self.tree, bg="#111827") for _ in range(4)]

        y_scroll = ttk.Scrollbar(tree_box, orient="vertical", command=self.tree_yview)
        self.tree.configure(yscrollcommand=lambda first, last: self.tree_yscroll_set(y_scroll, first, last))
        self.tree.pack(side="left", fill="both", expand=True)
        y_scroll.pack(side="right", fill="y")

        self.tree.bind("<Button-1>", self.select_clicked_cell)
        self.tree.bind("<Double-1>", self.start_cell_edit)
        self.tree.bind("<Control-c>", self.copy_selected_values)
        self.tree.bind("<Control-C>", self.copy_selected_values)
        self.tree.bind("<Control-x>", self.cut_selected_values)
        self.tree.bind("<Control-X>", self.cut_selected_values)
        self.tree.bind("<Control-v>", self.paste_from_clipboard)
        self.tree.bind("<Control-V>", self.paste_from_clipboard)
        self.tree.bind("<Delete>", self.clear_selected_values)
        self.tree.bind("<BackSpace>", self.clear_selected_values)
        self.tree.bind("<Key>", self.start_typing_in_selected_cell)
        self.tree.bind("<Up>", lambda _event: self.move_active_cell(-1, 0))
        self.tree.bind("<Down>", lambda _event: self.move_active_cell(1, 0))
        self.tree.bind("<Left>", lambda _event: self.move_active_cell(0, -1))
        self.tree.bind("<Right>", lambda _event: self.move_active_cell(0, 1))
        self.tree.bind("<Return>", lambda _event: self.start_cell_edit())
        self.tree.bind("<Tab>", lambda _event: self.move_active_cell(0, 1))
        self.tree.bind("<Shift-Tab>", lambda _event: self.move_active_cell(0, -1))
        self.tree.bind("<MouseWheel>", self.tree_mousewheel)
        self.tree.bind("<Button-4>", self.tree_mousewheel)
        self.tree.bind("<Button-5>", self.tree_mousewheel)
        self.tree.bind("<Configure>", lambda _event: self.draw_active_cell_border())
        self.tree.bind("<Enter>", self.disable_main_mousewheel)
        self.tree.bind("<Leave>", self.enable_main_mousewheel)
        self.register_drop_target(self.tree)
        self.register_drop_target(table_frame)
        self.update_time_table_labels()

    def update_time_table_labels(self):
        if self.source_mode_var.get() == "webcam":
            self.table_frame.config(text=" 4. 송출 경과시간별 수위값 표 ")
            self.tree.heading("time", text="경과시간")
        else:
            self.table_frame.config(text=" 4. 영상 시간별 수위값 표 ")
            self.tree.heading("time", text="영상 시간")

    def update_main_scroll_region(self, _event=None):
        self.main_canvas.configure(scrollregion=self.main_canvas.bbox("all"))

    def resize_main_canvas_window(self, event):
        self.main_canvas.itemconfigure(self.main_canvas_window, width=event.width)

    def enable_main_mousewheel(self, _event=None):
        self.root.bind_all("<MouseWheel>", self.on_main_mousewheel)
        self.root.bind_all("<Button-4>", self.on_main_mousewheel)
        self.root.bind_all("<Button-5>", self.on_main_mousewheel)

    def disable_main_mousewheel(self, _event=None):
        self.root.unbind_all("<MouseWheel>")
        self.root.unbind_all("<Button-4>")
        self.root.unbind_all("<Button-5>")

    def on_main_mousewheel(self, event):
        if getattr(event, "num", None) == 4:
            delta = -3
        elif getattr(event, "num", None) == 5:
            delta = 3
        else:
            delta = int(-1 * (event.delta / 120)) if event.delta else 0
        if delta:
            self.main_canvas.yview_scroll(delta, "units")

    def set_server_status(self, text, color, ready=False):
        if self.closing:
            return
        self.lbl_server.config(text=text, fg=color)
        if not self.is_streaming:
            self.btn_action.config(state="normal" if ready else "disabled")

    def start_internal_server(self):
        if not os.path.isfile(MEDIAMTX_PATH):
            messagebox.showerror("실행 오류", "mediamtx.exe를 찾을 수 없습니다.")
            self.set_server_status("RTSP 서버: 실행 파일 없음", "#dc2626")
            return
        if not os.path.isfile(MEDIAMTX_CONFIG_PATH):
            messagebox.showerror("실행 오류", "mediamtx.yml 설정 파일을 찾을 수 없습니다.")
            self.set_server_status("RTSP 서버: 설정 파일 없음", "#dc2626")
            return
        threading.Thread(target=self._start_server_worker, daemon=True).start()

    def _start_server_worker(self):
        if self._port_is_open("127.0.0.1", 8554):
            self.root.after(
                0,
                lambda: self.set_server_status(
                    "RTSP 서버: 기존 서버 사용 중 (127.0.0.1:8554)", "#16a34a", True
                ),
            )
            return

        try:
            self.server_log_file = open(
                MEDIAMTX_LOG_PATH, "w", encoding="utf-8", errors="replace"
            )
            self.server_process = subprocess.Popen(
                [MEDIAMTX_PATH, MEDIAMTX_CONFIG_PATH],
                cwd=CURRENT_DIR,
                stdout=self.server_log_file,
                stderr=subprocess.STDOUT,
                creationflags=CREATE_NO_WINDOW,
            )

            deadline = time.time() + 5
            while time.time() < deadline:
                if self.server_process.poll() is not None:
                    raise RuntimeError(self._read_server_log())
                if self._port_is_open("127.0.0.1", 8554):
                    self.root.after(
                        0,
                        lambda: self.set_server_status(
                            "RTSP 서버: 정상 실행 중 (127.0.0.1:8554)", "#16a34a", True
                        ),
                    )
                    return
                time.sleep(0.1)
            raise RuntimeError("5초 안에 8554 포트가 열리지 않았습니다.")
        except Exception as exc:
            detail = str(exc).strip() or "알 수 없는 서버 실행 오류"
            self.root.after(0, lambda: self._show_server_error(detail))

    def _show_server_error(self, detail):
        self.set_server_status("RTSP 서버: 실행 실패", "#dc2626")
        messagebox.showerror(
            "RTSP 서버 실행 실패",
            f"내장 RTSP 서버를 시작하지 못했습니다.\n\n{detail}",
        )

    @staticmethod
    def _port_is_open(host, port):
        try:
            with socket.create_connection((host, port), timeout=0.25):
                return True
        except OSError:
            return False

    def _read_server_log(self):
        if self.server_log_file:
            self.server_log_file.flush()
        try:
            with open(MEDIAMTX_LOG_PATH, "r", encoding="utf-8", errors="replace") as log:
                return log.read()[-3000:]
        except OSError:
            return "MediaMTX가 시작 직후 종료되었습니다."

    def update_source_mode_ui(self):
        if hasattr(self, "table_frame"):
            self.update_time_table_labels()

        if self.source_mode_var.get() == "webcam":
            self.btn_browse.pack_forget()
            self.combo_webcam.pack(side="right", padx=5)
            self.btn_refresh_webcam.pack(side="right", padx=5)
            self.preview_frame.pack(fill="x", pady=(6, 0), after=self.source_frame)
            self.lbl_file.config(text="웹캠 검색 중...", fg="#555555")
            self.lbl_duration.config(text="웹캠 라이브 입력", fg="#555555")
            current_device = self.webcam_var.get().strip()
            if self.webcam_devices and current_device not in self.webcam_devices:
                current_device = self.webcam_devices[0]
                self.webcam_var.set(current_device)

            if not self.webcam_devices or not current_device:
                self.refresh_webcam_devices(show_error=False)
            elif current_device:
                self.lbl_file.config(text=f"웹캠: {current_device}", fg="black")
                self.start_webcam_preview()
        else:
            self.stop_webcam_preview()
            self.preview_frame.pack_forget()
            self.combo_webcam.pack_forget()
            self.btn_refresh_webcam.pack_forget()
            self.btn_browse.pack(side="right", padx=5)
            if self.video_path:
                self.lbl_file.config(text=os.path.basename(self.video_path), fg="black")
            else:
                self.lbl_file.config(text="선택된 파일이 없습니다.", fg="gray")
            if self.video_duration > 0:
                self.lbl_duration.config(
                    text=f"영상 길이: {self.format_time(self.video_duration)}", fg="#555555"
                )
            else:
                self.lbl_duration.config(text="", fg="#555555")

    def refresh_webcam_devices(self, show_error=True):
        if not os.path.isfile(FFMPEG_PATH):
            if show_error:
                messagebox.showerror("실행 오류", "ffmpeg.exe를 찾을 수 없습니다.")
            return

        try:
            result = subprocess.run(
                [FFMPEG_PATH, "-hide_banner", "-list_devices", "true", "-f", "dshow", "-i", "dummy"],
                cwd=CURRENT_DIR,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.PIPE,
                text=True,
                encoding="utf-8",
                errors="replace",
                creationflags=CREATE_NO_WINDOW,
                timeout=10,
            )
        except Exception as exc:
            if show_error:
                messagebox.showerror("웹캠 확인 실패", str(exc))
            return

        devices = self.parse_dshow_video_devices(result.stderr)
        self.webcam_devices = devices
        self.combo_webcam["values"] = devices
        if devices and self.webcam_var.get() not in devices:
            self.webcam_var.set(devices[0])
            self.lbl_file.config(text=f"웹캠: {devices[0]}", fg="black")
            self.start_webcam_preview()
        elif devices and self.webcam_var.get() in devices:
            self.lbl_file.config(text=f"웹캠: {self.webcam_var.get()}", fg="black")
            self.start_webcam_preview()
        elif not devices:
            self.stop_webcam_preview()
            self.webcam_var.set("")
            self.lbl_file.config(text="웹캠을 찾지 못했습니다.", fg="#dc2626")
            self.lbl_preview.config(
                image="",
                text="웹캠을 찾지 못했습니다.",
                bg="black",
                fg="white",
                width=44,
                height=10,
            )
            self.lbl_preview_status.config(text="장치 연결 상태를 확인한 뒤 웹캠 새로고침을 눌러 주세요.", fg="#dc2626")
            if show_error:
                messagebox.showwarning("웹캠 없음", "FFmpeg에서 인식한 웹캠 장치가 없습니다.")
        elif self.source_mode_var.get() == "webcam":
            self.start_webcam_preview()

    @staticmethod
    def parse_dshow_video_devices(text):
        devices = []
        in_video_section = False
        for line in text.splitlines():
            lowered = line.lower()

            typed_match = re.search(r'"([^"]+)"\s*\((video|audio)\)', line, re.IGNORECASE)
            if typed_match:
                name, device_type = typed_match.groups()
                if device_type.lower() == "video" and name not in devices:
                    devices.append(name)
                continue

            if "directshow video devices" in lowered:
                in_video_section = True
                continue
            if "directshow audio devices" in lowered:
                in_video_section = False
                continue
            if not in_video_section:
                continue
            match = re.search(r'"([^"]+)"', line)
            if match:
                name = match.group(1)
                if name.startswith("@device_"):
                    continue
                if name not in devices:
                    devices.append(name)
        return devices

    def on_webcam_selected(self, _event=None):
        if self.webcam_var.get():
            self.lbl_file.config(text=f"웹캠: {self.webcam_var.get()}", fg="black")
            self.start_webcam_preview()

    def start_webcam_preview(self):
        if self.closing or self.is_streaming or self.source_mode_var.get() != "webcam":
            return

        device_name = self.webcam_var.get().strip()
        if not device_name:
            return
        if not os.path.isfile(FFMPEG_PATH):
            self.lbl_preview_status.config(text="ffmpeg.exe를 찾을 수 없어 미리보기를 실행할 수 없습니다.", fg="#dc2626")
            return

        if (
            self.preview_process
            and self.preview_process.poll() is None
            and getattr(self.preview_process, "device_name", None) == device_name
        ):
            return

        self.stop_webcam_preview(clear_image=False)
        self.lbl_preview.config(
            image="",
            text="웹캠 미리보기 연결 중...",
            bg="black",
            fg="white",
            width=44,
            height=10,
        )
        self.lbl_preview_status.config(
            text="선택한 웹캠 화면을 확인하는 중입니다.\n송출 시작 시 미리보기는 자동 중지됩니다.",
            fg="#555555",
        )

        command = [
            FFMPEG_PATH,
            "-hide_banner",
            "-loglevel",
            "error",
            "-f",
            "dshow",
            "-i",
            f"video={device_name}",
            "-an",
            "-vf",
            "fps=5,scale=360:-1",
            "-f",
            "image2pipe",
            "-vcodec",
            "ppm",
            "pipe:1",
        ]

        try:
            process = subprocess.Popen(
                command,
                cwd=CURRENT_DIR,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                creationflags=CREATE_NO_WINDOW,
            )
        except Exception as exc:
            self.lbl_preview_status.config(text=f"웹캠 미리보기 시작 실패: {exc}", fg="#dc2626")
            return

        process.device_name = device_name
        self.preview_process = process
        self.preview_thread = threading.Thread(
            target=self._run_webcam_preview,
            args=(process, device_name),
            daemon=True,
        )
        self.preview_thread.start()
        self.schedule_preview_poll()

    def stop_webcam_preview(self, clear_image=True):
        process = self.preview_process
        self.preview_process = None
        if process and process.poll() is None:
            try:
                process.terminate()
                process.wait(timeout=1)
            except (subprocess.TimeoutExpired, OSError):
                try:
                    process.kill()
                    process.wait(timeout=1)
                except OSError:
                    pass

        if clear_image and hasattr(self, "lbl_preview"):
            self.preview_image = None
            self.lbl_preview.config(
                image="",
                text="웹캠을 선택하면 여기에 미리보기가 표시됩니다.",
                bg="black",
                fg="white",
                width=44,
                height=10,
            )
            self.lbl_preview_status.config(text="", fg="#555555")

    def _run_webcam_preview(self, process, device_name):
        last_error = ""
        try:
            while not self.closing and process.poll() is None and self.preview_process is process:
                frame = self._read_ppm_frame(process.stdout)
                if frame is None:
                    break
                self.preview_queue.put(("frame", process, frame))
        except Exception as exc:
            last_error = str(exc)
        finally:
            if process.poll() is None:
                try:
                    process.terminate()
                    process.wait(timeout=1)
                except (subprocess.TimeoutExpired, OSError):
                    try:
                        process.kill()
                        process.wait(timeout=1)
                    except OSError:
                        pass

            if process.stderr:
                try:
                    error_bytes = process.stderr.read()
                    if error_bytes:
                        last_error = error_bytes.decode("utf-8", errors="replace").strip() or last_error
                except OSError:
                    pass

            if not self.closing and self.preview_process is process:
                self.preview_process = None
                self.preview_queue.put(("stopped", process, device_name, last_error))

    def schedule_preview_poll(self):
        if self.closing or self.preview_polling:
            return
        self.preview_polling = True
        self.root.after(50, self.poll_preview_queue)

    def poll_preview_queue(self):
        self.preview_polling = False
        while True:
            try:
                event = self.preview_queue.get_nowait()
            except queue.Empty:
                break

            event_type = event[0]
            if event_type == "frame":
                _, process, frame = event
                self._show_preview_frame(process, frame)
            elif event_type == "stopped":
                _, _process, device_name, error_text = event
                self._show_preview_stopped(device_name, error_text)

        if not self.closing and (self.preview_process or not self.preview_queue.empty()):
            self.schedule_preview_poll()

    def _show_preview_frame(self, process, ppm_data):
        if self.closing or self.preview_process is not process:
            return
        try:
            self.preview_image = tk.PhotoImage(data=ppm_data, format="PPM")
            self.lbl_preview.config(
                image=self.preview_image,
                text="",
                width=self.preview_image.width(),
                height=self.preview_image.height(),
            )
            self.lbl_preview_status.config(
                text=f"미리보기 정상 표시 중\n장치: {self.webcam_var.get()}",
                fg="#16a34a",
            )
        except tk.TclError as exc:
            self.lbl_preview_status.config(text=f"미리보기 표시 실패: {exc}", fg="#dc2626")

    def _show_preview_stopped(self, device_name, error_text):
        if self.closing or self.source_mode_var.get() != "webcam" or self.webcam_var.get() != device_name:
            return
        detail = error_text.strip()
        if detail:
            detail = detail[-500:]
            message = f"미리보기 실패\n{detail}"
        else:
            message = "미리보기가 중지되었습니다."
        self.preview_image = None
        self.lbl_preview.config(
            image="",
            text="웹캠 미리보기 없음",
            bg="black",
            fg="white",
            width=44,
            height=10,
        )
        self.lbl_preview_status.config(text=message, fg="#dc2626" if error_text else "#555555")

    @classmethod
    def _read_ppm_frame(cls, stream):
        magic = cls._read_ppm_token(stream)
        if not magic:
            return None
        if magic != b"P6":
            return None
        width = cls._read_ppm_token(stream)
        height = cls._read_ppm_token(stream)
        max_value = cls._read_ppm_token(stream)
        if not width or not height or max_value != b"255":
            return None

        frame_size = int(width) * int(height) * 3
        pixels = stream.read(frame_size)
        if len(pixels) != frame_size:
            return None
        return b"P6\n" + width + b" " + height + b"\n255\n" + pixels

    @staticmethod
    def _read_ppm_token(stream):
        token = bytearray()
        while True:
            char = stream.read(1)
            if not char:
                return None
            if char == b"#":
                stream.readline()
                continue
            if char in b" \t\r\n":
                continue
            token.extend(char)
            break

        while True:
            char = stream.read(1)
            if not char:
                break
            if char in b" \t\r\n":
                break
            token.extend(char)
        return bytes(token)

    def browse_file(self):
        selected_file = filedialog.askopenfilename(
            title="송출할 영상 선택",
            filetypes=[
                ("동영상 파일", "*.mp4 *.avi *.mkv *.mov *.wmv *.webm *.ts"),
                ("모든 파일", "*.*"),
            ],
        )
        if not selected_file:
            return

        self.video_path = selected_file
        self.lbl_file.config(text=os.path.basename(selected_file), fg="black")
        self.video_duration = self.get_video_duration(selected_file)
        if self.video_duration <= 0:
            self.lbl_duration.config(text="길이 확인 실패", fg="#dc2626")
            self.schedule_rows = []
            self.refresh_tree()
            return

        self.lbl_duration.config(
            text=f"영상 길이: {self.format_time(self.video_duration)}", fg="#555555"
        )
        self.rebuild_schedule_from_ui()

    def register_drop_target(self, widget):
        if DND_FILES is None or not hasattr(widget, "drop_target_register"):
            return
        try:
            widget.drop_target_register(DND_FILES)
            widget.dnd_bind("<<Drop>>", self.handle_file_drop)
        except tk.TclError:
            pass

    def handle_file_drop(self, event):
        try:
            paths = self.root.tk.splitlist(event.data)
        except tk.TclError:
            paths = [event.data]
        if not paths:
            return
        self.import_level_values_from_path(paths[0])

    def import_level_file(self):
        selected_file = filedialog.askopenfilename(
            title="수위값 파일 선택",
            filetypes=[
                ("엑셀/CSV 파일", "*.xlsx *.xlsm *.xltx *.xltm *.xls *.csv *.tsv *.txt"),
                ("Excel 파일", "*.xlsx *.xlsm *.xltx *.xltm *.xls"),
                ("CSV/텍스트 파일", "*.csv *.tsv *.txt"),
                ("모든 파일", "*.*"),
            ],
        )
        if selected_file:
            self.import_level_values_from_path(selected_file)

    def import_level_values_from_path(self, path):
        if self.is_streaming:
            messagebox.showwarning("불러오기 제한", "송출 중에는 수위값 파일을 불러올 수 없습니다.")
            return

        try:
            values = self.read_level_values_file(path)
        except Exception as exc:
            messagebox.showerror("수위값 불러오기 실패", str(exc))
            return

        if self.skip_file_first_row_var.get() and values:
            values = values[1:]

        values = [value for value in values if str(value).strip() != ""]
        if not values:
            messagebox.showwarning("수위값 없음", "파일에서 불러올 수위값을 찾지 못했습니다.")
            return

        inserted = self.fill_values_from_selected_row(values)
        messagebox.showinfo(
            "수위값 불러오기",
            f"{os.path.basename(path)} 파일에서 {inserted}개 값을 입력했습니다.",
        )

    def read_level_values_file(self, path):
        ext = os.path.splitext(path)[1].lower()
        if ext in (".xlsx", ".xlsm", ".xltx", ".xltm"):
            return self.read_xlsx_values(path)
        if ext == ".xls":
            return self.read_xls_values(path)
        if ext in (".csv", ".tsv", ".txt"):
            return self.read_csv_values(path)
        raise ValueError("지원하지 않는 파일 형식입니다. xlsx, xls, csv, tsv, txt 파일을 사용해 주세요.")

    def read_xlsx_values(self, path):
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            values = []
            for row in sheet.iter_rows(values_only=True):
                value = self.pick_value_from_row(row)
                if value != "":
                    values.append(value)
            return values
        finally:
            workbook.close()

    def read_xls_values(self, path):
        workbook = xlrd.open_workbook(path)
        sheet = workbook.sheet_by_index(0)
        values = []
        for row_index in range(sheet.nrows):
            value = self.pick_value_from_row(sheet.row_values(row_index))
            if value != "":
                values.append(value)
        return values

    def read_csv_values(self, path):
        last_error = None
        for encoding in ("utf-8-sig", "cp949", "euc-kr", "utf-8"):
            try:
                with open(path, "r", encoding=encoding, newline="") as file:
                    sample = file.read(4096)
                    file.seek(0)
                    try:
                        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
                    except csv.Error:
                        dialect = csv.excel_tab if path.lower().endswith(".tsv") else csv.excel
                    reader = csv.reader(file, dialect)
                    values = []
                    for row in reader:
                        value = self.pick_value_from_row(row)
                        if value != "":
                            values.append(value)
                    return values
            except UnicodeDecodeError as exc:
                last_error = exc
        raise ValueError(f"파일 인코딩을 읽지 못했습니다: {last_error}")

    @staticmethod
    def pick_value_from_row(row):
        cells = []
        for cell in row:
            if cell is None:
                continue
            text = str(cell).strip()
            if text != "":
                cells.append(text)
        return cells[-1] if cells else ""

    def fill_values_from_selected_row(self, values):
        if not self.schedule_rows:
            if self.source_mode_var.get() == "webcam":
                row_count = len(values) + (1 if self.skip_table_first_row_var.get() else 0)
                if not self.rebuild_schedule_by_count(row_count):
                    return 0
            else:
                messagebox.showwarning("표 없음", "먼저 영상 파일을 선택해서 시간표를 만들어 주세요.")
                return 0

        if not self.schedule_rows:
            return 0

        selected = self.tree.selection()
        start_index = int(selected[0]) if selected else 0
        if self.skip_table_first_row_var.get() and start_index == 0:
            start_index = 1
        inserted = 0

        for offset, value in enumerate(values):
            index = start_index + offset
            if index >= len(self.schedule_rows):
                break
            value_text = str(value).strip()
            self.schedule_rows[index]["value"] = value_text
            self.tree.set(str(index), "value", value_text)
            inserted += 1

        if inserted:
            item = str(start_index)
            self.tree.selection_set(item)
            self.tree.focus(item)
            self.tree.see(item)
        return inserted

    def load_recent_settings(self):
        try:
            with open(RECENT_SETTINGS_PATH, "r", encoding="utf-8") as file:
                data = json.load(file)
        except (OSError, json.JSONDecodeError):
            data = []

        self.recent_settings = [
            path for path in data if isinstance(path, str) and os.path.isfile(path)
        ][:10]
        self.refresh_recent_settings_combo()

    def save_recent_settings(self):
        try:
            with open(RECENT_SETTINGS_PATH, "w", encoding="utf-8") as file:
                json.dump(self.recent_settings[:10], file, ensure_ascii=False, indent=2)
        except OSError:
            pass

    def refresh_recent_settings_combo(self):
        if not hasattr(self, "combo_recent_settings"):
            return
        labels = [self.make_recent_label(path) for path in self.recent_settings]
        self.combo_recent_settings["values"] = labels
        if labels and self.recent_settings_var.get() not in labels:
            self.recent_settings_var.set(labels[0])
        elif not labels:
            self.recent_settings_var.set("")

    @staticmethod
    def make_recent_label(path):
        return os.path.basename(path) or path

    def remember_setting_path(self, path):
        path = os.path.abspath(path)
        self.recent_settings = [item for item in self.recent_settings if os.path.abspath(item) != path]
        self.recent_settings.insert(0, path)
        self.recent_settings = self.recent_settings[:10]
        self.save_recent_settings()
        self.refresh_recent_settings_combo()
        self.recent_settings_var.set(self.make_recent_label(path))

    def save_settings_as(self):
        self.finish_cell_edit(save=True)
        selected_file = filedialog.asksaveasfilename(
            title="설정 저장",
            defaultextension=".json",
            filetypes=[("RTSPStreamer 설정", "*.json"), ("모든 파일", "*.*")],
            initialfile="rtspstreamer_settings.json",
        )
        if not selected_file:
            return

        data = self.collect_settings()
        try:
            with open(selected_file, "w", encoding="utf-8") as file:
                json.dump(data, file, ensure_ascii=False, indent=2)
        except OSError as exc:
            messagebox.showerror("설정 저장 실패", str(exc))
            return

        self.remember_setting_path(selected_file)
        messagebox.showinfo("설정 저장", "설정을 저장했습니다.")

    def load_settings_from_file(self):
        selected_file = filedialog.askopenfilename(
            title="설정 불러오기",
            filetypes=[("RTSPStreamer 설정", "*.json"), ("모든 파일", "*.*")],
        )
        if selected_file:
            self.apply_settings_file(selected_file)

    def load_recent_setting(self, _event=None):
        selected_index = self.combo_recent_settings.current()
        if selected_index < 0 or selected_index >= len(self.recent_settings):
            return
        self.apply_settings_file(self.recent_settings[selected_index])

    def collect_settings(self):
        source_mode = self.source_mode_var.get()
        return {
            "version": 1,
            "source_mode": source_mode,
            "video_path": self.video_path if source_mode == "file" else "",
            "webcam_device": self.webcam_var.get() if source_mode == "webcam" else "",
            "rtsp_url": self.entry_url.get().strip(),
            "db_enabled": self.db_enabled_var.get(),
            "db_name": self.entry_db_name.get().strip(),
            "db_user": self.entry_db_user.get().strip(),
            "db_password": self.entry_db_password.get(),
            "db_host": self.entry_db_host.get().strip(),
            "db_port": self.entry_db_port.get().strip(),
            "schema": self.entry_schema.get().strip(),
            "table": self.entry_table.get().strip(),
            "time_column": self.entry_time_col.get().strip(),
            "value_column": self.entry_value_col.get().strip(),
            "interval_value": self.interval_value_var.get(),
            "interval_unit": self.interval_unit_var.get(),
            "video_duration": self.video_duration,
            "schedule_rows": self.schedule_rows,
        }

    def apply_settings_file(self, path):
        try:
            with open(path, "r", encoding="utf-8") as file:
                data = json.load(file)
        except (OSError, json.JSONDecodeError) as exc:
            messagebox.showerror("설정 불러오기 실패", str(exc))
            return

        try:
            self.apply_settings(data)
        except Exception as exc:
            messagebox.showerror("설정 적용 실패", str(exc))
            return

        self.remember_setting_path(path)

    def apply_settings(self, data):
        self.finish_cell_edit(save=False)
        source_mode = data.get("source_mode", "file")
        self.source_mode_var.set(source_mode if source_mode in ("file", "webcam") else "file")
        self.video_path = str(data.get("video_path", "") or "")
        self.webcam_var.set(str(data.get("webcam_device", "") or ""))
        if self.video_path:
            self.lbl_file.config(text=os.path.basename(self.video_path), fg="black")
        else:
            self.lbl_file.config(text="선택된 파일이 없습니다.", fg="gray")

        if self.video_path and os.path.isfile(self.video_path):
            self.video_duration = self.get_video_duration(self.video_path)
        else:
            self.video_duration = float(data.get("video_duration", 0) or 0)

        if self.video_duration > 0:
            self.lbl_duration.config(
                text=f"영상 길이: {self.format_time(self.video_duration)}", fg="#555555"
            )
        elif self.video_path:
            self.lbl_duration.config(text="영상 파일 없음/길이 확인 실패", fg="#dc2626")
        else:
            self.lbl_duration.config(text="", fg="#555555")

        self.update_source_mode_ui()
        if self.source_mode_var.get() == "webcam" and self.webcam_var.get():
            self.lbl_file.config(text=f"웹캠: {self.webcam_var.get()}", fg="black")

        self.set_entry_value(self.entry_url, data.get("rtsp_url", DEFAULT_RTSP_URL))
        self.db_enabled_var.set(bool(data.get("db_enabled", True)))
        self.set_entry_value(self.entry_db_name, data.get("db_name", data.get("database", "MFLOW")))
        self.set_entry_value(self.entry_db_user, data.get("db_user", data.get("user", "postgres")))
        self.set_entry_value(self.entry_db_password, data.get("db_password", data.get("password", "")))
        self.set_entry_value(self.entry_db_host, data.get("db_host", data.get("host", "localhost")))
        self.set_entry_value(self.entry_db_port, data.get("db_port", data.get("port", "5432")))
        self.set_entry_value(self.entry_schema, data.get("schema", "MEASURE"))
        self.set_entry_value(self.entry_table, data.get("table", "water_level_log"))
        self.set_entry_value(self.entry_time_col, data.get("time_column", "created_at"))
        self.set_entry_value(self.entry_value_col, data.get("value_column", "level_value"))

        self.interval_value_var.set(str(data.get("interval_value", "5") or "5"))
        interval_unit = data.get("interval_unit", "초")
        self.interval_unit_var.set(interval_unit if interval_unit in ("초", "분") else "초")

        loaded_rows = data.get("schedule_rows", [])
        if isinstance(loaded_rows, list) and loaded_rows:
            self.schedule_rows = []
            for row in loaded_rows:
                if not isinstance(row, dict):
                    continue
                offset = float(row.get("offset", 0) or 0)
                time_text = str(row.get("time") or self.format_time(offset))
                value = str(row.get("value", "") or "")
                self.schedule_rows.append({"offset": offset, "time": time_text, "value": value})
            self.refresh_tree()
        elif self.video_duration > 0:
            self.rebuild_schedule_from_ui()
        else:
            self.schedule_rows = []
            self.refresh_tree()

    @staticmethod
    def set_entry_value(entry, value):
        entry.delete(0, tk.END)
        entry.insert(0, str(value or ""))

    def get_video_duration(self, path):
        try:
            result = subprocess.run(
                [FFMPEG_PATH, "-hide_banner", "-i", path],
                cwd=CURRENT_DIR,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.PIPE,
                text=True,
                encoding="utf-8",
                errors="replace",
                creationflags=CREATE_NO_WINDOW,
                timeout=15,
            )
        except Exception:
            return 0.0

        match = re.search(r"Duration:\s*(\d+):(\d+):(\d+(?:\.\d+)?)", result.stderr)
        if not match:
            return 0.0
        hours, minutes, seconds = match.groups()
        return int(hours) * 3600 + int(minutes) * 60 + float(seconds)

    def get_interval_seconds_from_ui(self):
        try:
            interval_value = int(self.interval_value_var.get())
            if interval_value <= 0:
                raise ValueError
        except ValueError:
            messagebox.showwarning("간격 확인", "간격은 1 이상의 숫자로 입력해 주세요.")
            return None

        return interval_value * 60 if self.interval_unit_var.get() == "분" else interval_value

    def rebuild_schedule_by_count(self, row_count, interval_seconds=None):
        if row_count <= 0:
            return False
        if interval_seconds is None:
            interval_seconds = self.get_interval_seconds_from_ui()
        if interval_seconds is None:
            return False

        old_values = [row.get("value", "") for row in self.schedule_rows]
        rows = []
        for index in range(row_count):
            offset = index * interval_seconds
            rows.append(
                {
                    "offset": float(offset),
                    "time": self.format_time(offset),
                    "value": old_values[index] if index < len(old_values) else "",
                }
            )
        self.schedule_rows = rows
        self.refresh_tree()
        return True

    def rebuild_schedule_from_ui(self):
        interval_seconds = self.get_interval_seconds_from_ui()
        if interval_seconds is None:
            return

        if self.source_mode_var.get() == "webcam":
            if not self.schedule_rows:
                messagebox.showwarning(
                    "표 없음",
                    "웹캠은 전체 길이를 알 수 없어서 빈 표를 자동 생성할 수 없습니다.\n"
                    "수위파일을 불러오면 값 개수에 맞춰 송출 경과시간표가 자동 생성됩니다.",
                )
                return
            self.rebuild_schedule_by_count(len(self.schedule_rows), interval_seconds)
            return

        if self.video_duration <= 0:
            messagebox.showwarning("영상 확인", "먼저 영상 파일을 선택해 주세요.")
            return

        old_values = {row["time"]: row.get("value", "") for row in self.schedule_rows}

        rows = []
        offset = 0
        while offset < self.video_duration:
            time_text = self.format_time(offset)
            rows.append({"offset": float(offset), "time": time_text, "value": old_values.get(time_text, "")})
            offset += interval_seconds

        last_text = self.format_time(self.video_duration)
        if not rows or rows[-1]["time"] != last_text:
            rows.append(
                {
                    "offset": float(self.video_duration),
                    "time": last_text,
                    "value": old_values.get(last_text, ""),
                }
            )

        self.schedule_rows = rows
        self.refresh_tree()

    def add_schedule_row(self):
        interval_seconds = self.get_interval_seconds_from_ui()
        if interval_seconds is None:
            return

        if self.schedule_rows:
            last_offset = float(self.schedule_rows[-1].get("offset", 0) or 0)
            next_offset = last_offset + interval_seconds
        else:
            next_offset = 0.0

        self.schedule_rows.append(
            {
                "offset": float(next_offset),
                "time": self.format_time(next_offset),
                "value": "",
            }
        )
        self.refresh_tree()
        item = str(len(self.schedule_rows) - 1)
        self.tree.selection_set(item)
        self.tree.focus(item)
        self.tree.see(item)

    def delete_selected_rows(self):
        self.finish_cell_edit(save=False)
        selected = self.tree.selection()
        if not selected:
            focused = self.tree.focus()
            selected = (focused,) if focused else ()
        if not selected:
            return

        indexes = set()
        for item in selected:
            try:
                indexes.add(int(item))
            except ValueError:
                continue

        self.schedule_rows = [
            row for index, row in enumerate(self.schedule_rows) if index not in indexes
        ]
        self.refresh_tree()

        if self.schedule_rows:
            next_index = min(indexes) if indexes else 0
            next_index = min(next_index, len(self.schedule_rows) - 1)
            item = str(next_index)
            self.tree.selection_set(item)
            self.tree.focus(item)
            self.tree.see(item)

    def refresh_tree(self):
        self.finish_cell_edit(save=True)
        for item in self.tree.get_children():
            self.tree.delete(item)
        for index, row in enumerate(self.schedule_rows):
            self.tree.insert("", "end", iid=str(index), values=(row["time"], row.get("value", "")))
        if self.schedule_rows:
            if not self.active_cell_item or int(self.active_cell_item) >= len(self.schedule_rows):
                self.active_cell_item = "0"
            self.set_active_cell(self.active_cell_item, self.active_cell_column)
        else:
            self.active_cell_item = ""
            self.hide_active_cell_border()

    def clear_values(self):
        selected = self.tree.selection()
        if selected:
            self.clear_selected_values()
            return
        for row in self.schedule_rows:
            row["value"] = ""
        self.refresh_tree()

    @staticmethod
    def format_time(seconds):
        total = max(0, int(round(seconds)))
        hours = total // 3600
        minutes = (total % 3600) // 60
        secs = total % 60
        if hours:
            return f"{hours:02d}:{minutes:02d}:{secs:02d}"
        return f"{minutes:02d}:{secs:02d}"

    def tree_yview(self, *args):
        self.tree.yview(*args)
        self.root.after_idle(self.draw_active_cell_border)

    def tree_yscroll_set(self, scrollbar, first, last):
        scrollbar.set(first, last)
        self.root.after_idle(self.draw_active_cell_border)

    def tree_mousewheel(self, event):
        if getattr(event, "num", None) == 4:
            delta = -3
        elif getattr(event, "num", None) == 5:
            delta = 3
        else:
            delta = int(-1 * (event.delta / 120)) if event.delta else 0
        if delta:
            self.tree.yview_scroll(delta, "units")
            self.root.after_idle(self.draw_active_cell_border)
        return "break"

    @staticmethod
    def tree_column_from_id(column_id):
        return "time" if column_id == "#1" else "value"

    @staticmethod
    def tree_column_to_id(column_name):
        return "#1" if column_name == "time" else "#2"

    def set_active_cell(self, item, column_name):
        if not item:
            self.hide_active_cell_border()
            return
        if column_name not in ("time", "value"):
            column_name = "value"
        self.active_cell_item = str(item)
        self.active_cell_column = column_name
        self.tree.focus(str(item))
        self.tree.selection_set(str(item))
        self.tree.focus_set()
        self.draw_active_cell_border()

    def draw_active_cell_border(self):
        if not self.active_cell_item:
            self.hide_active_cell_border()
            return
        column_id = self.tree_column_to_id(self.active_cell_column)
        bbox = self.tree.bbox(self.active_cell_item, column_id)
        if not bbox:
            self.hide_active_cell_border()
            return
        x, y, width, height = bbox
        thickness = 2
        top, right, bottom, left = self.cell_border_parts
        top.place(x=x, y=y, width=width, height=thickness)
        bottom.place(x=x, y=y + height - thickness, width=width, height=thickness)
        left.place(x=x, y=y, width=thickness, height=height)
        right.place(x=x + width - thickness, y=y, width=thickness, height=height)

    def hide_active_cell_border(self):
        for part in self.cell_border_parts:
            part.place_forget()

    def select_clicked_cell(self, event):
        region = self.tree.identify("region", event.x, event.y)
        if region != "cell":
            return
        item = self.tree.identify_row(event.y)
        column_id = self.tree.identify_column(event.x)
        if not item or column_id not in ("#1", "#2"):
            return
        self.finish_cell_edit(save=True)
        self.set_active_cell(item, self.tree_column_from_id(column_id))
        return "break"

    def start_cell_edit(self, event=None):
        if event is not None:
            region = self.tree.identify("region", event.x, event.y)
            column_id = self.tree.identify_column(event.x)
            item = self.tree.identify_row(event.y)
            if region != "cell" or column_id not in ("#1", "#2") or not item:
                return "break"
            self.set_active_cell(item, self.tree_column_from_id(column_id))

        if not self.active_cell_item:
            return "break"
        self.open_cell_editor(self.active_cell_item, self.active_cell_column)
        return "break"

    def open_cell_editor(self, item, column_name, initial_text=None, append=False):
        self.finish_cell_edit(save=True)
        column_id = self.tree_column_to_id(column_name)
        bbox = self.tree.bbox(item, column_id)
        if not bbox:
            self.tree.see(item)
            self.root.update_idletasks()
            bbox = self.tree.bbox(item, column_id)
        if not bbox:
            return

        x, y, width, height = bbox
        value = self.tree.set(item, column_name)
        self.edit_entry = tk.Entry(self.tree)
        if initial_text is None:
            self.edit_entry.insert(0, value)
        elif append:
            self.edit_entry.insert(0, value + initial_text)
        else:
            self.edit_entry.insert(0, initial_text)
        self.edit_entry.select_range(0, tk.END)
        if initial_text is not None:
            self.edit_entry.icursor(tk.END)
        self.edit_entry.focus_set()
        self.edit_entry.place(x=x, y=y, width=width, height=height)
        self.edit_entry.item_id = item
        self.edit_entry.column_name = column_name
        self.edit_entry.bind("<Return>", lambda _event: self.move_from_editor(1))
        self.edit_entry.bind("<Down>", lambda _event: self.move_from_editor(1))
        self.edit_entry.bind("<Up>", lambda _event: self.move_from_editor(-1))
        self.edit_entry.bind("<Tab>", lambda _event: self.move_from_editor(0, 1))
        self.edit_entry.bind("<Shift-Tab>", lambda _event: self.move_from_editor(0, -1))
        self.edit_entry.bind("<Escape>", lambda _event: self.finish_cell_edit(save=False))
        self.edit_entry.bind("<FocusOut>", lambda _event: self.finish_cell_edit(save=True))

    def start_typing_in_selected_cell(self, event):
        if event.keysym in {
            "Up",
            "Down",
            "Left",
            "Right",
            "Home",
            "End",
            "Prior",
            "Next",
            "Tab",
            "Return",
            "Escape",
        }:
            return
        if event.state & 0x0004:
            return
        if event.keysym in {"Delete", "BackSpace"}:
            return
        if not event.char or ord(event.char) < 32:
            return

        item = self.active_cell_item or self.tree.focus()
        if not item:
            return "break"

        self.open_cell_editor(item, self.active_cell_column, initial_text=event.char)
        return "break"

    def finish_cell_edit(self, save):
        if not self.edit_entry:
            return
        entry = self.edit_entry
        self.edit_entry = None
        if save and hasattr(entry, "item_id"):
            item = entry.item_id
            column_name = getattr(entry, "column_name", "value")
            value = entry.get().strip()
            index = int(item)
            if 0 <= index < len(self.schedule_rows):
                if column_name == "time":
                    seconds = self.parse_time_text(value)
                    if seconds is None:
                        messagebox.showwarning(
                            "시간 형식 확인",
                            "시간은 00:10, 01:30, 01:02:03 또는 초 숫자 형식으로 입력해 주세요.",
                        )
                    else:
                        normalized = self.format_time(seconds)
                        self.schedule_rows[index]["offset"] = float(seconds)
                        self.schedule_rows[index]["time"] = normalized
                        if item in self.tree.get_children():
                            self.tree.set(item, "time", normalized)
                else:
                    self.schedule_rows[index]["value"] = value
                    if item in self.tree.get_children():
                        self.tree.set(item, "value", value)
        entry.destroy()
        self.root.after_idle(self.draw_active_cell_border)

    @staticmethod
    def parse_time_text(text):
        text = str(text).strip()
        if not text:
            return None
        if re.fullmatch(r"\d+(?:\.\d+)?", text):
            return max(0, float(text))
        parts = text.split(":")
        if len(parts) not in (2, 3):
            return None
        try:
            numbers = [int(part) for part in parts]
        except ValueError:
            return None
        if any(number < 0 for number in numbers):
            return None
        if len(numbers) == 2:
            minutes, seconds = numbers
            if seconds >= 60:
                return None
            return float(minutes * 60 + seconds)
        hours, minutes, seconds = numbers
        if minutes >= 60 or seconds >= 60:
            return None
        return float(hours * 3600 + minutes * 60 + seconds)

    def move_from_editor(self, row_delta=0, col_delta=0):
        item = self.edit_entry.item_id if self.edit_entry and hasattr(self.edit_entry, "item_id") else self.tree.focus()
        column_name = getattr(self.edit_entry, "column_name", self.active_cell_column) if self.edit_entry else self.active_cell_column
        self.finish_cell_edit(save=True)
        self.set_active_cell(item, column_name)
        self.move_active_cell(row_delta, col_delta)
        return "break"

    def move_active_cell(self, row_delta=0, col_delta=0):
        children = self.tree.get_children()
        if not children:
            return "break"

        current = self.active_cell_item or self.tree.focus() or children[0]
        try:
            current_index = children.index(current)
        except ValueError:
            current_index = 0

        next_index = min(max(current_index + row_delta, 0), len(children) - 1)
        next_item = children[next_index]
        columns = ["time", "value"]
        try:
            column_index = columns.index(self.active_cell_column)
        except ValueError:
            column_index = 1
        next_column = columns[min(max(column_index + col_delta, 0), len(columns) - 1)]
        self.set_active_cell(next_item, next_column)
        self.tree.see(next_item)
        self.root.after_idle(self.draw_active_cell_border)
        return "break"

    def paste_from_clipboard(self, _event=None):
        self.finish_cell_edit(save=True)
        try:
            text = self.root.clipboard_get()
        except tk.TclError:
            return "break"

        lines = [line for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n") if line != ""]
        if not lines:
            return "break"

        start_index = int(self.active_cell_item or self.tree.focus() or 0)
        start_column = self.active_cell_column
        column_order = ["time", "value"]
        start_col_index = column_order.index(start_column) if start_column in column_order else 1

        for row_offset, line in enumerate(lines):
            index = start_index + row_offset
            if index >= len(self.schedule_rows):
                break
            cells = line.split("\t")
            for col_offset, cell in enumerate(cells):
                col_index = start_col_index + col_offset
                if col_index >= len(column_order):
                    break
                column_name = column_order[col_index]
                value = cell.strip()
                if column_name == "time":
                    seconds = self.parse_time_text(value)
                    if seconds is None:
                        continue
                    normalized = self.format_time(seconds)
                    self.schedule_rows[index]["offset"] = float(seconds)
                    self.schedule_rows[index]["time"] = normalized
                    self.tree.set(str(index), "time", normalized)
                else:
                    self.schedule_rows[index]["value"] = value
                    self.tree.set(str(index), "value", value)

        return "break"

    def get_selected_items_for_clipboard(self):
        selected = [self.active_cell_item] if self.active_cell_item else []
        if not selected:
            focused = self.tree.focus()
            selected = [focused] if focused else []
        return sorted(selected, key=lambda item: int(item))

    def copy_selected_values(self, _event=None):
        self.finish_cell_edit(save=True)
        item = self.active_cell_item or self.tree.focus()
        if not item:
            return "break"

        try:
            index = int(item)
        except ValueError:
            return "break"
        if not (0 <= index < len(self.schedule_rows)):
            return "break"

        if self.active_cell_column == "time":
            value = str(self.schedule_rows[index].get("time", ""))
        else:
            value = str(self.schedule_rows[index].get("value", ""))

        self.root.clipboard_clear()
        self.root.clipboard_append(value)
        return "break"

    def cut_selected_values(self, _event=None):
        self.copy_selected_values()
        self.clear_selected_values()
        return "break"

    def clear_selected_values(self, _event=None):
        self.finish_cell_edit(save=False)
        item = self.active_cell_item or self.tree.focus()
        if not item:
            return "break"
        try:
            index = int(item)
        except ValueError:
            return "break"

        if 0 <= index < len(self.schedule_rows):
            if self.active_cell_column == "time":
                self.schedule_rows[index]["offset"] = 0.0
                self.schedule_rows[index]["time"] = "00:00"
                self.tree.set(item, "time", "00:00")
            else:
                self.schedule_rows[index]["value"] = ""
                self.tree.set(item, "value", "")
        self.draw_active_cell_border()

        return "break"

    def toggle_streaming(self):
        if self.is_streaming:
            self.stop_streaming()
        else:
            self.start_streaming()

    def start_streaming(self):
        self.finish_cell_edit(save=True)
        source_mode = self.source_mode_var.get()
        if source_mode == "file" and not os.path.isfile(self.video_path):
            messagebox.showwarning("파일 확인", "먼저 송출할 영상 파일을 선택해 주세요.")
            return
        if source_mode == "webcam" and not self.webcam_var.get().strip():
            messagebox.showwarning("웹캠 확인", "송출할 웹캠을 선택해 주세요.")
            return
        if not os.path.isfile(FFMPEG_PATH):
            messagebox.showerror("실행 오류", "ffmpeg.exe를 찾을 수 없습니다.")
            return

        rtsp_url = self.entry_url.get().strip()
        parsed = urlparse(rtsp_url)
        if parsed.scheme.lower() not in ("rtsp", "rtsps") or not parsed.hostname:
            messagebox.showwarning(
                "주소 확인", "올바른 RTSP 주소를 입력해 주세요.\n예: rtsp://127.0.0.1:8554/live/stream"
            )
            return

        if self.db_enabled_var.get():
            try:
                self.prepare_database()
            except Exception as exc:
                messagebox.showerror("DB 설정 오류", str(exc))
                return

        if source_mode == "webcam":
            self.stop_webcam_preview(clear_image=False)
            self.lbl_preview_status.config(
                text="송출 중에는 카메라 충돌 방지를 위해 미리보기를 잠시 중지합니다.",
                fg="#d97706",
            )

        self.is_streaming = True
        self.btn_action.config(text="연결 중...", bg="#d97706", state="disabled")
        self.btn_browse.config(state="disabled")
        self.combo_webcam.config(state="disabled")
        self.btn_refresh_webcam.config(state="disabled")
        self.entry_url.config(state="disabled")
        self.lbl_status.config(text="상태: RTSP 서버 연결 중...", fg="#d97706")

        if self.db_enabled_var.get():
            self.db_thread = threading.Thread(target=self._run_db_scheduler, daemon=True)
            self.db_thread.start()

        threading.Thread(target=self._run_ffmpeg, args=(rtsp_url,), daemon=True).start()

    def validate_identifier(self, label, value, allow_empty=False):
        value = value.strip()
        if allow_empty and not value:
            return ""
        if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", value):
            raise ValueError(f"{label}은 영문/숫자/밑줄만 사용할 수 있고 숫자로 시작할 수 없습니다: {value}")
        return value

    @staticmethod
    def quote_identifier(value):
        return '"' + value.replace('"', '""') + '"'

    def get_db_config(self):
        db_name = self.entry_db_name.get().strip()
        db_user = self.entry_db_user.get().strip()
        db_password = self.entry_db_password.get()
        db_host = self.entry_db_host.get().strip()
        db_port = self.entry_db_port.get().strip()

        missing = []
        if not db_name:
            missing.append("DB 명")
        if not db_user:
            missing.append("DB 사용자 명")
        if not db_password:
            missing.append("패스워드")
        if not db_host:
            missing.append("접속주소(host)")
        if not db_port:
            missing.append("접속포트(port)")
        if missing:
            raise ValueError("필수 DB 정보를 입력해 주세요: " + ", ".join(missing))

        try:
            port_number = int(db_port)
        except ValueError as exc:
            raise ValueError("접속포트(port)는 숫자로 입력해 주세요.") from exc

        return {
            "db_name": db_name,
            "db_user": db_user,
            "db_password": db_password,
            "db_host": db_host,
            "db_port": port_number,
            "schema": self.validate_identifier("스키마명", self.entry_schema.get(), allow_empty=True),
            "table": self.validate_identifier("테이블명", self.entry_table.get()),
            "time_col": self.validate_identifier("현재시간 컬럼명", self.entry_time_col.get()),
            "value_col": self.validate_identifier("수위값 컬럼명", self.entry_value_col.get()),
        }

    def open_db_connection(self, config):
        return psycopg2.connect(
            dbname=config["db_name"],
            user=config["db_user"],
            password=config["db_password"],
            host=config["db_host"],
            port=config["db_port"],
            connect_timeout=5,
        )

    @staticmethod
    def table_identifier(config):
        if config.get("schema"):
            return sql.Identifier(config["schema"], config["table"])
        return sql.Identifier(config["table"])

    def test_db_connection(self):
        try:
            config = self.get_db_config()
            with self.open_db_connection(config) as conn:
                with conn.cursor() as cursor:
                    cursor.execute("SELECT 1")
                    cursor.fetchone()
                    if config.get("schema"):
                        cursor.execute(
                            """
                            SELECT 1
                            FROM information_schema.tables
                            WHERE table_schema = %s AND table_name = %s
                            """,
                            (config["schema"], config["table"]),
                        )
                    else:
                        cursor.execute(
                            """
                            SELECT 1
                            FROM information_schema.tables
                            WHERE table_name = %s
                            """,
                            (config["table"],),
                        )
                    table_exists = cursor.fetchone() is not None
        except Exception as exc:
            messagebox.showerror("DB 연결 실패", str(exc))
            return
        table_name = f"{config['schema']}.{config['table']}" if config.get("schema") else config["table"]
        suffix = f"\n대상 테이블 확인: {table_name}" if table_exists else f"\n대상 테이블 없음: {table_name}\n송출 시작 시 자동 생성합니다."
        messagebox.showinfo("DB 연결 성공", "PostgreSQL DB 연결에 성공했습니다." + suffix)

    def prepare_database(self):
        config = self.get_db_config()

        columns = [
            sql.SQL("{} TIMESTAMP NOT NULL").format(sql.Identifier(config["time_col"]))
        ]
        columns.append(sql.SQL("{} DOUBLE PRECISION").format(sql.Identifier(config["value_col"])))

        query = sql.SQL("CREATE TABLE IF NOT EXISTS {} ({})").format(
            self.table_identifier(config),
            sql.SQL(", ").join(columns),
        )

        with self.open_db_connection(config) as conn:
            with conn.cursor() as cursor:
                cursor.execute(query)
            conn.commit()

    def _run_db_scheduler(self):
        rows = [
            dict(row)
            for row in self.schedule_rows
            if str(row.get("value", "")).strip() != ""
        ]
        if not rows:
            self.root.after(
                0,
                lambda: self.lbl_db_status.config(
                    text="DB INSERT: 입력할 수위값 없음", fg="#d97706"
                ),
            )
            return

        config = self.get_db_config()
        start_mono = time.monotonic()
        cycle_index = 0
        inserted = set()
        insert_count = 0
        is_file_source = self.source_mode_var.get() == "file"
        duration = max(self.video_duration, 1.0)

        self.root.after(0, lambda: self.lbl_db_status.config(text="DB INSERT: 실행 중", fg="#16a34a"))

        while self.is_streaming and not self.closing:
            elapsed = time.monotonic() - start_mono
            cycle_start = cycle_index * duration

            for row_index, row in enumerate(rows):
                key = (cycle_index, row_index)
                target = cycle_start + float(row["offset"])
                if key not in inserted and elapsed >= target:
                    try:
                        self.insert_db_row(config, row)
                        insert_count += 1
                        inserted.add(key)
                        self.root.after(
                            0,
                            lambda count=insert_count: self.lbl_db_status.config(
                                text=f"DB INSERT: {count}건 입력", fg="#16a34a"
                            ),
                        )
                    except Exception as exc:
                        self.root.after(
                            0,
                            lambda err=str(exc): self.lbl_db_status.config(
                                text=f"DB INSERT 오류: {err[:80]}", fg="#dc2626"
                            ),
                        )

            if not is_file_source and len(inserted) >= len(rows):
                self.root.after(
                    0,
                    lambda count=insert_count: self.lbl_db_status.config(
                        text=f"DB INSERT: {count}건 입력 완료", fg="#16a34a"
                    ),
                )
                break

            if is_file_source and elapsed >= (cycle_index + 1) * duration + 0.2:
                cycle_index += 1
                inserted.clear()

            time.sleep(0.1)

    def insert_db_row(self, config, row):
        columns = [sql.Identifier(config["time_col"])]
        values = [datetime.now()]

        columns.append(sql.Identifier(config["value_col"]))
        raw_value = str(row.get("value", "")).strip()
        try:
            value = float(raw_value)
        except ValueError:
            raise ValueError(f"수위값은 숫자여야 합니다: {raw_value}")
        values.append(value)

        query = sql.SQL("INSERT INTO {} ({}) VALUES ({})").format(
            self.table_identifier(config),
            sql.SQL(", ").join(columns),
            sql.SQL(", ").join(sql.Placeholder() * len(values)),
        )

        with self.open_db_connection(config) as conn:
            with conn.cursor() as cursor:
                cursor.execute(query, values)
            conn.commit()

    def _run_ffmpeg(self, rtsp_url):
        if self.source_mode_var.get() == "webcam":
            attempts = [self._build_ffmpeg_command(rtsp_url, repair_timestamps=False)]
        else:
            attempts = [
                self._build_ffmpeg_command(rtsp_url, repair_timestamps=False),
                self._build_ffmpeg_command(rtsp_url, repair_timestamps=True),
            ]
        last_error = ""

        try:
            for attempt_number, command in enumerate(attempts):
                if not self.is_streaming or self.closing:
                    return
                if attempt_number == 1:
                    self.root.after(0, self._mark_repairing)

                return_code, error_text = self._execute_ffmpeg(command)
                last_error = error_text

                if return_code == 0 or not self.is_streaming or self.closing:
                    return
                if attempt_number == 0 and self._is_timestamp_error(error_text):
                    continue
                break

            if self.is_streaming and not self.closing:
                detail = last_error.strip() or "FFmpeg가 오류 코드와 함께 종료되었습니다."
                self.root.after(0, lambda: self._show_stream_error(detail[-4000:]))
        except Exception as exc:
            if not self.closing:
                self.root.after(0, lambda: self._show_stream_error(str(exc)))
        finally:
            self.ffmpeg_process = None
            if not self.closing:
                self.root.after(0, self._reset_stream_ui)

    def _build_ffmpeg_command(self, rtsp_url, repair_timestamps):
        command = [
            FFMPEG_PATH,
            "-hide_banner",
            "-loglevel",
            "warning",
            "-re",
        ]

        if self.source_mode_var.get() == "webcam":
            command.extend(
                [
                    "-f",
                    "dshow",
                    "-i",
                    f"video={self.webcam_var.get()}",
                ]
            )
        else:
            command.extend(
                [
                    "-stream_loop",
                    "-1",
                ]
            )

        if repair_timestamps and self.source_mode_var.get() != "webcam":
            command.extend(
                [
                    "-ignore_editlist",
                    "1",
                    "-fflags",
                    "+genpts+discardcorrupt",
                    "-err_detect",
                    "ignore_err",
                ]
            )

        if self.source_mode_var.get() != "webcam":
            command.extend(["-i", self.video_path])

        command.extend(
            [
                "-map",
                "0:v:0",
                "-map",
                "0:a:0?",
                "-c:v",
                "libx264",
                "-preset",
                "veryfast",
                "-tune",
                "zerolatency",
                "-pix_fmt",
                "yuv420p",
                "-g",
                "60",
                "-keyint_min",
                "60",
                "-sc_threshold",
                "0",
                "-c:a",
                "aac",
                "-b:a",
                "128k",
                "-ar",
                "48000",
                "-ac",
                "2",
                "-f",
                "rtsp",
                "-rtsp_transport",
                "tcp",
                rtsp_url,
            ]
        )
        return command

    def _execute_ffmpeg(self, command):
        self.ffmpeg_process = subprocess.Popen(
            command,
            cwd=CURRENT_DIR,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.PIPE,
            creationflags=CREATE_NO_WINDOW,
        )

        time.sleep(1)
        if self.ffmpeg_process.poll() is None and self.is_streaming:
            self.root.after(0, self._mark_streaming)

        _, stderr = self.ffmpeg_process.communicate()
        return_code = self.ffmpeg_process.returncode
        error_text = stderr.decode("utf-8", errors="replace").strip()

        try:
            with open(FFMPEG_LOG_PATH, "w", encoding="utf-8", errors="replace") as log:
                log.write(error_text)
        except OSError:
            pass

        self.ffmpeg_process = None
        return return_code, error_text

    @staticmethod
    def _is_timestamp_error(error_text):
        lowered = error_text.lower()
        indicators = (
            "missing key frame",
            "cannot find an index entry",
            "edit list",
            "non-monotonous dts",
            "invalid, non monotonically increasing dts",
            "operation not permitted",
        )
        return any(indicator in lowered for indicator in indicators)

    def _mark_repairing(self):
        if not self.is_streaming or self.closing:
            return
        self.btn_action.config(text="복구 재시도 중...", bg="#d97706", state="disabled")
        self.lbl_status.config(
            text="상태: 영상 타임스탬프를 복구하여 다시 연결 중...", fg="#d97706"
        )

    def _mark_streaming(self):
        if not self.is_streaming or self.closing:
            return
        self.btn_action.config(text="송출 중지", bg="#dc2626", state="normal")
        if self.source_mode_var.get() == "webcam":
            self.lbl_status.config(text="상태: 송출 중 (웹캠 라이브 / 경과시간 기준)", fg="#16a34a")
        else:
            self.lbl_status.config(text="상태: 송출 중 (영상 반복 재생)", fg="#16a34a")

    def _show_stream_error(self, detail):
        messagebox.showerror(
            "RTSP 송출 실패",
            f"송출을 시작하지 못했거나 송출이 중단되었습니다.\n\n{detail}",
        )

    def stop_streaming(self):
        self.is_streaming = False
        process = self.ffmpeg_process
        if process and process.poll() is None:
            try:
                process.terminate()
                process.wait(timeout=2)
            except (subprocess.TimeoutExpired, OSError):
                try:
                    process.kill()
                    process.wait(timeout=1)
                except OSError:
                    pass
        self._reset_stream_ui()

    def _reset_stream_ui(self):
        if self.closing:
            return
        self.is_streaming = False
        self.btn_action.config(text="송출 시작", bg="#16a34a", state="normal")
        self.btn_browse.config(state="normal")
        self.combo_webcam.config(state="readonly")
        self.btn_refresh_webcam.config(state="normal")
        self.entry_url.config(state="normal")
        self.lbl_status.config(text="상태: 대기 중", fg="#2563eb")
        if self.db_enabled_var.get():
            self.lbl_db_status.config(text="DB INSERT: 대기", fg="#555555")
        if self.source_mode_var.get() == "webcam" and self.webcam_var.get():
            self.start_webcam_preview()

    def close(self):
        self.closing = True
        self.is_streaming = False
        self.stop_webcam_preview(clear_image=False)

        if self.ffmpeg_process and self.ffmpeg_process.poll() is None:
            try:
                self.ffmpeg_process.kill()
            except OSError:
                pass

        if self.server_process and self.server_process.poll() is None:
            try:
                self.server_process.terminate()
                self.server_process.wait(timeout=2)
            except (subprocess.TimeoutExpired, OSError):
                try:
                    self.server_process.kill()
                except OSError:
                    pass

        if self.server_log_file:
            self.server_log_file.close()
        self.root.destroy()


if __name__ == "__main__":
    root = TkinterDnD.Tk() if TkinterDnD is not None else tk.Tk()
    RTSPStreamerGUI(root)
    root.mainloop()
